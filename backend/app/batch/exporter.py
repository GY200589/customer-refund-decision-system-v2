"""导出挂起订单为 Excel"""
import io
import logging
from typing import Tuple

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from app.models import RefundCase

logger = logging.getLogger(__name__)

# 表头定义
HEADERS = [
    "case_id",
    "order_id",
    "金额(元)",
    "风险分",
    "舆情等级",
    "审批动作",
    "审批意见",
]

# 列宽（按字符数）
COL_WIDTHS = [36, 20, 12, 10, 12, 14, 30]

# 样式
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center")
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def export_suspended_cases(session) -> Tuple[bytes, int]:
    """查库导出 SUSPENDED 状态的工单为 Excel 字节流

    Args:
        session: SQLAlchemy 数据库会话

    Returns:
        Tuple[bytes, int]: (Excel 文件字节流, 记录数)

    Note:
        无挂起单时返回空表头 Excel（含表头，0 数据行）
    """
    cases = (
        session.query(RefundCase)
        .filter(RefundCase.status == "SUSPENDED")
        .order_by(RefundCase.updated_at.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "挂起审批工单"

    # 写表头
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER

    # 写数据行（审批动作列留空由主管填写）
    for row_idx, case in enumerate(cases, 2):
        ws.cell(row=row_idx, column=1, value=case.case_id).border = THIN_BORDER
        ws.cell(row=row_idx, column=2, value=case.order_id).border = THIN_BORDER
        ws.cell(row=row_idx, column=3, value=case.amount_cent / 100).border = THIN_BORDER
        ws.cell(row=row_idx, column=4, value=case.fraud_score).border = THIN_BORDER
        ws.cell(row=row_idx, column=5, value=case.risk_level or "N/A").border = THIN_BORDER
        # 第 6 列「审批动作」留空
        ws.cell(row=row_idx, column=6).border = THIN_BORDER
        # 第 7 列「审批意见」留空
        ws.cell(row=row_idx, column=7).border = THIN_BORDER

    # 列宽
    for col, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[chr(64 + col)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue(), len(cases)