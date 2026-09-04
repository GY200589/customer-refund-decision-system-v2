"""容器内文件解析脚本（沙箱执行体，独立运行）。

只依赖 python-docx / openpyxl，不 import 项目其他模块，因此可单独
COPY 进沙箱镜像执行。文件路径相对于容器工作目录（由挂载决定）。

用法：
  python container_runner.py read_word   <rel_path>
  python container_runner.py read_excel  <rel_path>
  python container_runner.py write_word  <rel_path>     # 文本内容从 stdin 读
  python container_runner.py write_excel <rel_path>     # 行数据 JSON 从 stdin 读

输出：单行 JSON，形如 {"ok": true, "result": ...} 或 {"ok": false, "error": ...}。
"""

import json
import os
import sys

# 在 read-only 容器中，设置临时目录为可写的 /tmp（由 --tmpfs 提供）
os.environ.setdefault("TMPDIR", "/tmp")
os.environ.setdefault("TEMP", "/tmp")
os.environ.setdefault("TMP", "/tmp")


def _emit(ok: bool, result=None, error: str = None) -> None:
    payload = {"ok": ok}
    if ok:
        payload["result"] = result
    else:
        payload["error"] = error or "unknown error"
    print(json.dumps(payload, ensure_ascii=False))
    sys.exit(0 if ok else 1)


def read_word(path: str) -> str:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    tables = []
    for table in doc.tables:
        for row in table.rows:
            tables.append(" | ".join(cell.text for cell in row.cells))
    content = "\n".join(paragraphs)
    if tables:
        content += "\n\n--- Tables ---\n" + "\n".join(tables)
    return content


def read_excel(path: str, sheet_name=None):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []

    headers = [str(h) if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        result.append({
            headers[i] if i < len(headers) else f"col_{i}":
                (cell if cell is not None else "")
            for i, cell in enumerate(row)
        })
    return result


def write_word(path: str) -> str:
    from docx import Document

    text = sys.stdin.read()
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(path)
    return path


def write_excel(path: str) -> str:
    import openpyxl

    data = json.loads(sys.stdin.read() or "[]")
    wb = openpyxl.Workbook()
    ws = wb.active
    if data:
        headers = list(data[0].keys())
        ws.append(headers)
        for row in data:
            ws.append([row.get(h, "") for h in headers])
    wb.save(path)
    return path


def main() -> None:
    if len(sys.argv) < 3:
        _emit(False, error="usage: container_runner.py <op> <rel_path>")
    op, path = sys.argv[1], sys.argv[2]
    try:
        if op == "read_word":
            result = read_word(path)
        elif op == "read_excel":
            result = read_excel(path)
        elif op == "write_word":
            result = write_word(path)
        elif op == "write_excel":
            result = write_excel(path)
        else:
            _emit(False, error=f"unknown op: {op}")
        _emit(True, result=result)
    except Exception as e:  # noqa: BLE001 —— 容器内异常统一归一为 JSON 返回
        _emit(False, error=str(e))


if __name__ == "__main__":
    main()
